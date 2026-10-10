from __future__ import annotations

import base64
import io
import json
import os
import random
import re
import sqlite3
import time
from copy import deepcopy
from datetime import date, datetime, timedelta
from html import escape
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests
import streamlit as st
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import altair as alt
except Exception:
    # Streamlit Cloud側の依存関係に一時的な不整合があっても、
    # 順位確認・計測・Excel出力は利用できるようにする。
    alt = None


APP_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv("RANK_TRACKER_DB", APP_DIR / "rank_history.db"))
CSV_PATH = APP_DIR / "rank_history.csv"
CONFIG_PATH = APP_DIR / "rank_config.json"
SERPER_SEARCH_ENDPOINT = "https://google.serper.dev/search"
SERPER_MAPS_ENDPOINT = "https://google.serper.dev/maps"
HISTORY_COLUMNS = ["checked_at", "target_url", "keyword", "rank", "matched_url", "provider"]
APP_VERSION = "2026.10.10-gbp-cid-fix-v3"

DEFAULT_CONFIG = {
    "sites": [
        {
            "id": "ponte_nene_jp",
            "label": "ponte-nene.jp",
            "url": "https://ponte-nene.jp/",
            "keywords": [
                "越谷 整体", "越谷 整骨院", "越谷 交通事故治療", "越谷 訪問マッサージ",
                "せんげん台 整体", "せんげん台 整骨院", "せんげん台 交通事故治療", "せんげん台 訪問マッサージ",
            ],
        },
        {
            "id": "ponte_aroma_jp",
            "label": "ponte-aroma.jp",
            "url": "https://ponte-aroma.jp/",
            "keywords": [
                "越谷 マッサージ", "越谷 エステ", "越谷 フェイシャル", "越谷 オイルマッサージ",
                "せんげん台 マッサージ", "せんげん台 エステ", "せんげん台 フェイシャル", "せんげん台 オイルマッサージ",
            ],
        },
        {
            "id": "ponte_nene_net",
            "label": "ponte-nene.net",
            "url": "https://ponte-nene.net/",
            "keywords": ["越谷 おすすめ"],
        },
    ],
    "gbp": {
        "profiles": [
            {
                "id": "gbp_ponte_nene",
                "label": "ぽんて鍼灸整骨院",
                "target": "gbp:5558918595392795634",
                "maps_url": "https://maps.app.goo.gl/uVZhgeL29M1kHYvYA?g_st=ic",
                "keywords": [
                    "越谷 整体", "越谷 整骨院", "越谷 交通事故治療", "越谷 訪問マッサージ",
                    "せんげん台 整体", "せんげん台 整骨院", "せんげん台 交通事故治療", "せんげん台 訪問マッサージ",
                ],
            },
            {
                "id": "gbp_ponte_aroma",
                "label": "ぽんてアロマサロン",
                "target": "gbp:120948687201481996",
                "maps_url": "https://maps.app.goo.gl/L4bCG7AZKwzkhpFx5?g_st=ic",
                "keywords": [
                    "越谷 マッサージ", "越谷 エステ", "越谷 フェイシャル", "越谷 オイルマッサージ",
                    "せんげん台 マッサージ", "せんげん台 エステ", "せんげん台 フェイシャル", "せんげん台 オイルマッサージ",
                ],
            },
        ]
    },
}


def secret_value(name: str) -> str:
    try:
        return str(st.secrets.get(name, ""))
    except Exception:
        return ""


def load_config() -> dict:
    config = deepcopy(DEFAULT_CONFIG)
    if not CONFIG_PATH.exists():
        return config
    try:
        loaded = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return config

    sites = loaded.get("sites")
    if isinstance(sites, list) and sites:
        config["sites"] = sites
    elif isinstance(sites, dict) and sites:
        converted = []
        for site_id, value in sites.items():
            if isinstance(value, dict):
                converted.append({"id": site_id, **value})
        if converted:
            config["sites"] = converted

    profiles = loaded.get("gbp", {}).get("profiles") if isinstance(loaded.get("gbp"), dict) else None
    if isinstance(profiles, list) and profiles:
        config["gbp"]["profiles"] = profiles
    elif isinstance(profiles, dict) and profiles:
        converted = []
        for profile_id, value in profiles.items():
            if isinstance(value, dict):
                converted.append({"id": profile_id, **value})
        if converted:
            config["gbp"]["profiles"] = converted
    return config


def save_keywords(item_type: str, item: dict, keywords: list[str]) -> tuple[bool, str, dict]:
    """変更をセッションへ保存し、可能ならサーバー上のrank_config.jsonにも反映する。"""
    session_config = st.session_state.get("editable_rank_config")
    if isinstance(session_config, dict):
        raw = deepcopy(session_config)
    else:
        try:
            if CONFIG_PATH.exists():
                raw = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            else:
                raw = deepcopy(DEFAULT_CONFIG)
        except (OSError, json.JSONDecodeError):
            raw = deepcopy(DEFAULT_CONFIG)

    if item_type == "site":
        container = raw.setdefault("sites", [])
        item_id = str(item.get("id") or "")
        target = site_target(item)
    else:
        gbp = raw.setdefault("gbp", {})
        if not isinstance(gbp, dict):
            gbp = {}
            raw["gbp"] = gbp
        container = gbp.setdefault("profiles", [])
        item_id = str(item.get("id") or "")
        target = profile_target(item)

    updated = False
    if isinstance(container, list):
        for saved_item in container:
            if not isinstance(saved_item, dict):
                continue
            saved_id = str(saved_item.get("id") or "")
            saved_target = site_target(saved_item) if item_type == "site" else profile_target(saved_item)
            if (item_id and saved_id == item_id) or (target and saved_target == target):
                saved_item["keywords"] = keywords
                updated = True
                break
        if not updated:
            new_item = deepcopy(item)
            new_item["keywords"] = keywords
            container.append(new_item)
            updated = True
    elif isinstance(container, dict):
        lookup_key = item_id
        if lookup_key and isinstance(container.get(lookup_key), dict):
            container[lookup_key]["keywords"] = keywords
            updated = True
        else:
            for saved_item in container.values():
                if not isinstance(saved_item, dict):
                    continue
                saved_target = site_target(saved_item) if item_type == "site" else profile_target(saved_item)
                if target and saved_target == target:
                    saved_item["keywords"] = keywords
                    updated = True
                    break

        if not updated:
            lookup_key = lookup_key or str(len(container) + 1)
            new_item = deepcopy(item)
            new_item["keywords"] = keywords
            container[lookup_key] = new_item
            updated = True

    if not updated:
        return False, "設定ファイル内で対象を確認できませんでした。", raw

    # Streamlit Cloudのファイル保存可否にかかわらず、現在の画面では保持する。
    st.session_state["editable_rank_config"] = deepcopy(raw)

    try:
        temporary_path = CONFIG_PATH.with_suffix(".json.tmp")
        temporary_path.write_text(json.dumps(raw, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        temporary_path.replace(CONFIG_PATH)
    except OSError as exc:
        return (
            True,
            "変更内容をこの画面に保存しました。サーバーへ直接保存できないため、下の設定ファイルをGitHubへアップロードしてください。",
            raw,
        )
    return True, "キーワードを更新しました。", raw


def config_download_html(config_data: dict | None = None) -> str:
    if isinstance(config_data, dict):
        payload = (json.dumps(config_data, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    else:
        try:
            payload = CONFIG_PATH.read_bytes() if CONFIG_PATH.exists() else (
                json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2) + "\n"
            ).encode("utf-8")
        except OSError:
            payload = (json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    encoded = base64.b64encode(payload).decode("ascii")
    return (
        '<a class="config-download" target="_blank" rel="noopener" '
        'download="rank_config.json" '
        f'href="data:application/json;base64,{encoded}">⚙️ 更新した設定ファイルをダウンロード</a>'
    )


def site_label(site: dict) -> str:
    return str(site.get("label") or site.get("name") or site.get("tab") or site.get("domain") or "Webサイト")


def site_target(site: dict) -> str:
    return str(site.get("url") or site.get("target_url") or site.get("target") or "")


def profile_label(profile: dict) -> str:
    return str(profile.get("label") or profile.get("name") or "GBP店舗")


def profile_target(profile: dict) -> str:
    target = str(profile.get("target") or profile.get("target_url") or "")
    if target:
        return target if target.startswith("gbp:") else f"gbp:{normalize_cid(target)}"
    cid = normalize_cid(profile.get("cid") or profile.get("data_cid"))
    return f"gbp:{cid}" if cid else ""


def init_db() -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """CREATE TABLE IF NOT EXISTS rankings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                checked_at TEXT NOT NULL,
                target_url TEXT NOT NULL,
                keyword TEXT NOT NULL,
                rank INTEGER,
                matched_url TEXT,
                provider TEXT NOT NULL
            )"""
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_rankings_lookup "
            "ON rankings(target_url, keyword, checked_at)"
        )


def normalize_url(value: str) -> str:
    value = str(value or "").strip()
    if value and not value.startswith(("http://", "https://")):
        value = "https://" + value
    parsed = urlparse(value)
    host = parsed.netloc.lower().removeprefix("www.")
    path = parsed.path.rstrip("/") or "/"
    return f"{host}{path}"


def target_matches(stored_target: str, requested_target: str) -> bool:
    stored = str(stored_target or "").strip()
    requested = str(requested_target or "").strip()
    if requested.startswith("gbp:"):
        return stored == requested
    return normalize_url(stored).split("/", 1)[0] == normalize_url(requested).split("/", 1)[0]


def is_web_match(result_url: str, target_url: str, match_mode: str) -> bool:
    result = normalize_url(result_url)
    target = normalize_url(target_url)
    if match_mode == "ドメイン全体":
        return result.split("/", 1)[0] == target.split("/", 1)[0]
    return result == target


def serper_web_rank(keyword: str, target_url: str, api_key: str, match_mode: str) -> tuple[int | None, str | None]:
    headers = {"X-API-KEY": api_key, "Content-Type": "application/json"}
    payload = {"q": keyword, "gl": "jp", "hl": "ja", "num": 100}
    response = requests.post(SERPER_SEARCH_ENDPOINT, headers=headers, json=payload, timeout=30)
    response.raise_for_status()
    data = response.json()
    if data.get("message"):
        raise RuntimeError(data["message"])
    for fallback_position, item in enumerate(data.get("organic", []), start=1):
        link = item.get("link", "")
        if link and is_web_match(link, target_url, match_mode):
            return int(item.get("position", fallback_position)), link
    return None, None


def normalize_cid(value: object) -> str:
    """GBPのCIDを比較可能な数字だけの形式へそろえる。

    アプリ内部では ``gbp:123...``、Serperの応答では ``123...`` や
    ``cid:123...``、URL内では ``?cid=123...`` になることがある。
    """
    text = str(value or "").strip()
    if not text:
        return ""

    lowered = text.lower()
    for prefix in ("gbp:", "cid:"):
        if lowered.startswith(prefix):
            text = text[len(prefix):].strip()
            lowered = text.lower()

    url_match = re.search(r"(?:[?&]|^)cid=(\d+)", text, flags=re.IGNORECASE)
    if url_match:
        return url_match.group(1)

    # CIDは数字。余計な空白や引用符が含まれた応答にも対応する。
    numeric_match = re.fullmatch(r"[^0-9]*(\d{10,})[^0-9]*", text)
    return numeric_match.group(1) if numeric_match else text


def serper_gbp_rank(keyword: str, target: str, api_key: str) -> tuple[int | None, str | None]:
    target_cid = normalize_cid(target)
    if not target_cid:
        raise ValueError("GBPのCIDが設定されていません。rank_config.jsonを確認してください。")
    headers = {"X-API-KEY": api_key, "Content-Type": "application/json"}
    payload = {"q": keyword, "gl": "jp", "hl": "ja", "num": 100}
    response = requests.post(SERPER_MAPS_ENDPOINT, headers=headers, json=payload, timeout=30)
    response.raise_for_status()
    data = response.json()
    if data.get("message"):
        raise RuntimeError(data["message"])
    places = data.get("places") or data.get("localResults") or []
    for fallback_position, item in enumerate(places, start=1):
        cid_candidates = (
            item.get("cid"),
            item.get("dataCid"),
            item.get("data_cid"),
            item.get("link"),
        )
        item_cids = {normalize_cid(candidate) for candidate in cid_candidates if candidate}
        if target_cid in item_cids:
            matched = item.get("link") or item.get("website") or item.get("title") or target
            return int(item.get("position", fallback_position)), str(matched)
    return None, None


def demo_rank(keyword: str, target: str) -> tuple[int | None, str | None]:
    seed = sum(ord(char) for char in keyword + target) + date.today().toordinal()
    rng = random.Random(seed)
    return rng.randint(1, 65), "デモデータ"


def save_result(target: str, keyword: str, rank: int | None, matched_url: str | None, provider: str) -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT INTO rankings(checked_at, target_url, keyword, rank, matched_url, provider) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (datetime.now().isoformat(timespec="seconds"), target, keyword, rank, matched_url, provider),
        )


def empty_history() -> pd.DataFrame:
    return pd.DataFrame(columns=HISTORY_COLUMNS)


def read_database_history() -> pd.DataFrame:
    if not DB_PATH.exists():
        return empty_history()
    try:
        with sqlite3.connect(DB_PATH) as conn:
            return pd.read_sql_query(
                "SELECT checked_at, target_url, keyword, rank, matched_url, provider FROM rankings",
                conn,
            )
    except (sqlite3.Error, pd.errors.DatabaseError):
        return empty_history()


def read_csv_history() -> pd.DataFrame:
    if not CSV_PATH.exists():
        return empty_history()
    try:
        frame = pd.read_csv(CSV_PATH)
    except (OSError, pd.errors.ParserError, UnicodeDecodeError):
        return empty_history()
    if "target" in frame.columns and "target_url" not in frame.columns:
        frame = frame.rename(columns={"target": "target_url"})
    for column in HISTORY_COLUMNS:
        if column not in frame.columns:
            frame[column] = None
    return frame[HISTORY_COLUMNS]


def load_history(target: str, days: int = 365) -> pd.DataFrame:
    frames = [frame for frame in (read_database_history(), read_csv_history()) if not frame.empty]
    if not frames:
        return empty_history()
    history = pd.concat(frames, ignore_index=True)
    history["checked_at"] = pd.to_datetime(history["checked_at"], errors="coerce", utc=True)
    history["rank"] = pd.to_numeric(history["rank"], errors="coerce")
    history = history.dropna(subset=["checked_at", "keyword", "target_url"])
    history = history[history["target_url"].map(lambda value: target_matches(value, target))]
    since = pd.Timestamp.now(tz="UTC") - pd.Timedelta(days=days)
    history = history[history["checked_at"] >= since]
    history = history.drop_duplicates(
        subset=["checked_at", "target_url", "keyword", "rank", "matched_url", "provider"],
        keep="last",
    )
    return history.sort_values(["checked_at", "keyword"]).reset_index(drop=True)


def latest_rows(history: pd.DataFrame, out_of_range: str) -> pd.DataFrame:
    latest = history.sort_values("checked_at").groupby("keyword", as_index=False).tail(1).copy()
    latest["順位"] = latest["rank"].apply(lambda value: f"{int(value)}位" if pd.notna(value) else out_of_range)
    latest["計測日時"] = latest["checked_at"].dt.strftime("%Y-%m-%d %H:%M")
    latest["検索ワード"] = latest["keyword"]
    latest["一致したURL・店舗"] = latest["matched_url"].fillna("—")
    return latest[["検索ワード", "順位", "計測日時", "一致したURL・店舗"]].sort_values("検索ワード")


def export_rows(history: pd.DataFrame, target_type: str, label: str, out_of_range: str) -> pd.DataFrame:
    columns = ["計測日時", "検索ワード", "順位", "状態", "一致したURL・店舗", "取得方法", "種別", "対象"]
    if history.empty:
        return pd.DataFrame(columns=columns)
    output = pd.DataFrame()
    output["計測日時"] = (
        pd.to_datetime(history["checked_at"], errors="coerce", utc=True)
        .dt.tz_convert("Asia/Tokyo")
        .dt.tz_localize(None)
    )
    output["検索ワード"] = history["keyword"].astype(str)
    output["順位"] = pd.to_numeric(history["rank"], errors="coerce")
    output["状態"] = output["順位"].apply(lambda value: "計測済み" if pd.notna(value) else out_of_range)
    output["一致したURL・店舗"] = history["matched_url"].fillna("").astype(str)
    output["取得方法"] = history["provider"].fillna("").astype(str)
    output["種別"] = target_type
    output["対象"] = label
    return output[columns].sort_values(["計測日時", "検索ワード"], ascending=[False, True])


def build_combined_excel(config: dict) -> tuple[bytes, dict[str, int]]:
    targets: list[tuple[str, str, str, str, str]] = []
    for index, site in enumerate(config.get("sites", [])[:3], start=1):
        label = site_label(site)
        targets.append((f"{index}_Web_{label}", "Web", label, site_target(site), "100位圏外"))
    for index, profile in enumerate(config.get("gbp", {}).get("profiles", [])[:2], start=4):
        label = profile_label(profile)
        targets.append((f"{index}_GBP_{label}", "GBP", label, profile_target(profile), "取得範囲外"))

    buffer = io.BytesIO()
    counts: dict[str, int] = {}
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, target_type, label, target, out_of_range in targets:
            safe_sheet_name = sheet_name[:31]
            rows = export_rows(load_history(target), target_type, label, out_of_range)
            rows.to_excel(writer, sheet_name=safe_sheet_name, index=False)
            counts[label] = len(rows)
            sheet = writer.book[safe_sheet_name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            sheet.row_dimensions[1].height = 26
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="3155D9" if target_type == "Web" else "0F766E")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            widths = [20, 28, 10, 16, 48, 16, 10, 24]
            for col_index, width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(col_index)].width = width
            for cell in sheet["A"][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm"
            for cell in sheet["C"][1:]:
                cell.number_format = '0"位"'
            if sheet.max_row >= 2:
                green = PatternFill("solid", fgColor="DCFCE7")
                yellow = PatternFill("solid", fgColor="FEF3C7")
                red = PatternFill("solid", fgColor="FEE2E2")
                sheet.conditional_formatting.add(f"C2:C{sheet.max_row}", CellIsRule(operator="lessThanOrEqual", formula=["10"], fill=green))
                sheet.conditional_formatting.add(f"C2:C{sheet.max_row}", CellIsRule(operator="between", formula=["11", "50"], fill=yellow))
                sheet.conditional_formatting.add(f"C2:C{sheet.max_row}", CellIsRule(operator="greaterThan", formula=["50"], fill=red))
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top")
        if not targets:
            pd.DataFrame({"メッセージ": ["設定対象がありません"]}).to_excel(writer, sheet_name="データなし", index=False)
    buffer.seek(0)
    return buffer.getvalue(), counts


def render_excel_download(config: dict) -> None:
    excel_bytes, counts = build_combined_excel(config)
    file_name = f"google_rank_history_5sites_{date.today().strftime('%Y%m%d')}.xlsx"
    encoded = base64.b64encode(excel_bytes).decode("ascii")
    total_rows = sum(counts.values())
    st.markdown(
        f"""
        <div class="download-card">
          <div class="download-title">5つの順位履歴をExcelで一括保存</div>
          <div class="download-note">Web 3サイト＋GBP 2店舗／全{total_rows:,}件／Excel内で5シートに分かれます。</div>
          <a class="excel-download" target="_blank" rel="noopener"
             download="{escape(file_name)}"
             href="data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64,{encoded}">
             📊 Excelを一括ダウンロード
          </a>
          <div class="iphone-note">iPhoneでは別画面で開きます。元のアプリは別タブに残るため、タブ切替で戻れます。</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def display_history(target: str, tab_key: str, out_of_range: str, max_rank: int) -> None:
    history = load_history(target)
    if history.empty:
        st.info("「分析開始」を押すと、ここに最新順位と履歴グラフが表示されます。")
        return
    valid = history.dropna(subset=["rank"])
    latest_valid = valid.sort_values("checked_at").groupby("keyword", as_index=False).tail(1)
    m1, m2, m3 = st.columns(3)
    m1.metric("登録キーワード", history["keyword"].nunique())
    m2.metric("10位以内", int((latest_valid["rank"] <= 10).sum()))
    m3.metric("最新の平均順位", f"{latest_valid['rank'].mean():.1f}位" if not latest_valid.empty else "—")
    st.subheader("最新の検索順位")
    st.dataframe(latest_rows(history, out_of_range), use_container_width=True, hide_index=True)
    st.subheader("過去1年間の順位変動")
    available = sorted(history["keyword"].dropna().unique())
    selected = st.multiselect("グラフに表示する検索ワード", available, default=available[:5], key=f"chart_{tab_key}")
    chart_data = history[history["keyword"].isin(selected)].dropna(subset=["rank"]).copy()
    if chart_data.empty:
        st.info("表示できる順位履歴がありません。")
        return
    if alt is None:
        st.warning("グラフ用ライブラリを読み込めないため、順位履歴を表で表示しています。")
        fallback = chart_data.pivot_table(index="checked_at", columns="keyword", values="rank", aggfunc="last")
        st.dataframe(fallback.sort_index(ascending=False), use_container_width=True)
        return
    upper = max(max_rank, int(chart_data["rank"].max()))
    chart = (
        alt.Chart(chart_data)
        .mark_line(point=True)
        .encode(
            x=alt.X("checked_at:T", title="計測日"),
            y=alt.Y("rank:Q", title="順位", scale=alt.Scale(reverse=True, domain=[1, upper])),
            color=alt.Color("keyword:N", title="検索ワード"),
            tooltip=[
                alt.Tooltip("checked_at:T", title="計測日時"),
                alt.Tooltip("keyword:N", title="検索ワード"),
                alt.Tooltip("rank:Q", title="順位"),
            ],
        )
        .properties(height=420)
        .interactive()
    )
    st.altair_chart(chart, use_container_width=True)


def run_measurement(target: str, keywords: list[str], target_type: str, api_key: str, demo_mode: bool, match_mode: str) -> None:
    if not keywords:
        st.error("検索ワードを1つ以上入力してください。")
        return
    if not demo_mode and not api_key:
        st.error("本番計測にはSerper APIキーが必要です。Secretsまたは左側の設定欄に登録してください。")
        return
    progress = st.progress(0, text="計測を開始しています…")
    errors: list[str] = []
    provider = "demo" if demo_mode else ("serper_maps" if target_type == "GBP" else "serper")
    for index, keyword in enumerate(keywords, start=1):
        try:
            if demo_mode:
                rank, matched = demo_rank(keyword, target)
                time.sleep(0.05)
            elif target_type == "GBP":
                rank, matched = serper_gbp_rank(keyword, target, api_key)
            else:
                rank, matched = serper_web_rank(keyword, target, api_key, match_mode)
            save_result(target, keyword, rank, matched, provider)
        except Exception as exc:
            errors.append(f"{keyword}: {exc}")
        progress.progress(index / len(keywords), text=f"{index}/{len(keywords)}件を計測中：{keyword}")
    progress.empty()
    if errors:
        st.error("一部の計測に失敗しました。\n\n" + "\n".join(errors))
    else:
        st.success(f"{len(keywords)}件の順位を保存しました。")


def render_web_tracker(site: dict, api_key: str, demo_mode: bool, match_mode: str) -> None:
    key = site.get("id", site_label(site))
    target_url = st.text_input("特定のWebページ", value=site_target(site), key=f"url_{key}")
    keyword_text = st.text_area(
        "検索ワードの一覧表（1行に1語）",
        value="\n".join(site.get("keywords", [])),
        height=190,
        key=f"keywords_{key}",
    )
    keywords = list(dict.fromkeys(line.strip() for line in keyword_text.splitlines() if line.strip()))
    action_col, save_col = st.columns(2)
    with action_col:
        run_clicked = st.button("🔍 分析開始", type="primary", use_container_width=True, key=f"run_{key}")
    with save_col:
        save_clicked = st.button("💾 キーワードを更新", use_container_width=True, key=f"save_{key}")
    if run_clicked:
        if not target_url.strip():
            st.error("WebページのURLを入力してください。")
        else:
            run_measurement(target_url.strip(), keywords, "Web", api_key, demo_mode, match_mode)
    if save_clicked:
        if not keywords:
            st.error("保存するキーワードを1つ以上入力してください。")
        else:
            success, message, updated_config = save_keywords("site", site, keywords)
            if success:
                st.success(f"{message} 現在は{len(keywords)}語です。")
                st.info("変更を消さず、毎週月曜日の自動計測にも反映するには、下のファイルをGitHubのrank_config.jsonと置き換えてください。")
                st.markdown(config_download_html(updated_config), unsafe_allow_html=True)
            else:
                st.error(message)
    display_history(target_url.strip(), key, "100位圏外", 100)


def render_gbp_tracker(profile: dict, api_key: str, demo_mode: bool, match_mode: str) -> None:
    key = profile.get("id", profile_label(profile))
    st.markdown(f"**対象店舗：{profile_label(profile)}**")
    maps_url = profile.get("maps_url") or profile.get("url")
    if maps_url:
        st.link_button("Googleマップで店舗を確認", maps_url, use_container_width=True)
    keyword_text = st.text_area(
        "検索ワードの一覧表（1行に1語）",
        value="\n".join(profile.get("keywords", [])),
        height=190,
        key=f"keywords_{key}",
    )
    keywords = list(dict.fromkeys(line.strip() for line in keyword_text.splitlines() if line.strip()))
    action_col, save_col = st.columns(2)
    with action_col:
        run_clicked = st.button("📍 GBP順位を分析", type="primary", use_container_width=True, key=f"run_{key}")
    with save_col:
        save_clicked = st.button("💾 キーワードを更新", use_container_width=True, key=f"save_{key}")
    if run_clicked:
        run_measurement(profile_target(profile), keywords, "GBP", api_key, demo_mode, match_mode)
    if save_clicked:
        if not keywords:
            st.error("保存するキーワードを1つ以上入力してください。")
        else:
            success, message, updated_config = save_keywords("gbp", profile, keywords)
            if success:
                st.success(f"{message} 現在は{len(keywords)}語です。")
                st.info("変更を消さず、毎週月曜日の自動計測にも反映するには、下のファイルをGitHubのrank_config.jsonと置き換えてください。")
                st.markdown(config_download_html(updated_config), unsafe_allow_html=True)
            else:
                st.error(message)
    display_history(profile_target(profile), key, "取得範囲外", 100)


st.set_page_config(page_title="Google順位チェッカー", page_icon="📈", layout="wide")
init_db()
config = load_config()
saved_api_key = secret_value("SERPER_API_KEY") or os.getenv("SERPER_API_KEY", "")

st.markdown(
    """<style>
    .block-container {max-width:1180px; padding-top:2rem; padding-bottom:4rem;}
    [data-testid="stMetric"] {background:#fff; border:1px solid #e6e9ef; padding:16px; border-radius:14px;}
    .hero {padding:1.5rem 1.6rem; border-radius:20px; color:white; background:linear-gradient(120deg,#3155d9,#11a9a0); margin-bottom:1.2rem;}
    .hero h1 {margin:0 0 .35rem 0; font-size:clamp(2rem,7vw,3.4rem); line-height:1.15;}
    .hero p {margin:0; opacity:.92; font-size:1.05rem;}
    .download-card {border:1px solid #dfe4ee; border-radius:18px; padding:1.15rem; margin:1rem 0 1.35rem; background:#f8fafc;}
    .download-title {font-weight:700; font-size:1.12rem; margin-bottom:.25rem; color:#1f2937;}
    .download-note,.iphone-note {font-size:.9rem; color:#64748b; margin-bottom:.8rem;}
    .iphone-note {margin:.7rem 0 0;}
    .excel-download {display:block; text-align:center; padding:.85rem 1rem; border-radius:12px; background:#16834a; color:white!important; text-decoration:none!important; font-weight:700; font-size:1.05rem;}
    .excel-download:hover {background:#116b3c;}
    .config-download {display:block; text-align:center; padding:.75rem .9rem; border-radius:10px; background:#475569; color:white!important; text-decoration:none!important; font-weight:700;}
    .config-download:hover {background:#334155;}
    @media (max-width:640px) {.block-container{padding-left:1rem;padding-right:1rem}.hero{padding:1.25rem}.hero p{font-size:.95rem}}
    </style>""",
    unsafe_allow_html=True,
)
st.markdown(
    '<div class="hero"><h1>Google順位<br>チェッカー</h1><p>Web検索とGoogleマップの掲載順位を記録し、1年間の変化を見える化します。</p></div>',
    unsafe_allow_html=True,
)
st.caption(f"アプリバージョン：{APP_VERSION}")

with st.sidebar:
    st.header("⚙️ 設定")
    if saved_api_key:
        st.success("Serper APIキーはSecretsから読み込み済みです。")
        api_key = saved_api_key
    else:
        api_key = st.text_input("Serper APIキー", type="password")
    demo_mode = st.toggle("デモモード", value=not bool(api_key), help="API通信を行わず画面を確認します。")
    match_mode = st.radio("URLの照合方法", ["入力ページのみ", "ドメイン全体"], index=1)
    st.caption("検索地域：日本 / 言語：日本語")

render_excel_download(config)

web_main, gbp_main = st.tabs(["🌐 Web検索順位", "📍 GBP順位"])
with web_main:
    web_sites = config.get("sites", [])[:3]
    web_tabs = st.tabs([site_label(site) or f"Web {index + 1}" for index, site in enumerate(web_sites)])
    for tab, site in zip(web_tabs, web_sites):
        with tab:
            render_web_tracker(site, api_key, demo_mode, match_mode)

with gbp_main:
    gbp_profiles = config.get("gbp", {}).get("profiles", [])[:2]
    gbp_tabs = st.tabs([profile_label(profile) or f"GBP {index + 1}" for index, profile in enumerate(gbp_profiles)])
    for tab, profile in zip(gbp_tabs, gbp_profiles):
        with tab:
            render_gbp_tracker(profile, api_key, demo_mode, match_mode)

if isinstance(st.session_state.get("editable_rank_config"), dict):
    st.markdown("### 💾 更新した検索ワードの保存")
    st.info("このファイルをGitHubのrank_config.jsonと置き換えると、再起動後も設定が残り、毎週月曜日の自動計測にも反映されます。")
    st.markdown(config_download_html(st.session_state["editable_rank_config"]), unsafe_allow_html=True)

with st.expander("ご利用前の注意"):
    st.markdown(
        """
        - 検索順位は地域・端末・時刻などで変わるため、実際の個人検索と差が出ることがあります。
        - Webの「100位圏外」およびGBPの「取得範囲外」は、取得結果内に対象が見つからなかった状態です。
        - Excelには直近1年間の履歴が、Web 3シート・GBP 2シートに分かれて保存されます。
        - キーワードを編集したら、同じタブの「キーワードを更新」を押してください。
        - 毎週月曜日の自動計測にも反映するには、表示されるrank_config.jsonをGitHubへアップロードしてください。
        - APIキーは順位取得にのみ利用し、順位履歴には保存しません。
        """
    )
